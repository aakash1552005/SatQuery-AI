import pptx
import sys

sys.stdout.reconfigure(encoding='utf-8')
prs = pptx.Presentation("PPT - SatSense .pptx")
slide = prs.slides[1]

print("=== SLIDE 2 DETAILED SHAPE MAP ===")
for i, s in enumerate(slide.shapes):
    txt = ""
    if s.has_text_frame:
        txt = " | ".join(p.text.strip() for p in s.text_frame.paragraphs if p.text.strip())
    print(f"[{i}] {s.name} | type={s.shape_type} | left={s.left.inches:.2f}, top={s.top.inches:.2f}, w={s.width.inches:.2f}, h={s.height.inches:.2f} | text='{txt[:60]}'")
