import pptx
import sys
from pathlib import Path

sys.stdout.reconfigure(encoding='utf-8')
target = Path("sih_presentation/SatQuery_AI_SIH2026_SatSense_V2.pptx")
if not target.exists():
    print(f"File not found: {target}")
    sys.exit(1)

prs = pptx.Presentation(str(target))
print(f"File: {target.name} ({target.stat().st_size / (1024**2):.2f} MB)")
print(f"Total Slides: {len(prs.slides)}")

for i, s in enumerate(prs.slides):
    title = "No Title"
    for shape in s.shapes:
        if shape.has_text_frame and shape.text_frame.text.strip():
            title = shape.text_frame.text.strip().split("\n")[0]
            break
    print(f"Slide {i+1}: {title[:60]}")
